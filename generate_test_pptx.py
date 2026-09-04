import os
import io
import zipfile
from pptx import Presentation
from pptx.util import Inches
from PIL import Image

OUTPUT_FOLDER = "Error Test PPTX"


def make_placeholder_image():
    image = Image.new("RGB", (100, 100), color="blue")
    image_stream = io.BytesIO()
    image.save(image_stream, format="PNG")
    image_stream.seek(0)
    return image_stream


def set_alt_text(picture, alt_text):
    picture._element.nvPicPr.cNvPr.set("descr", alt_text)


def clear_alt_text(picture):
    # add_picture() fills descr with the source filename by default,
    # so we have to clear it to truly simulate a picture with no alt text
    picture._element.nvPicPr.cNvPr.set("descr", "")


def save_presentation(prs, filename):
    output_path = os.path.join(OUTPUT_FOLDER, filename)
    prs.save(output_path)
    print(f"Created {filename}")


def make_control_no_issues():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    title_box.name = "Title 1"
    title_box.text_frame.text = "A Well-Formed Slide"

    picture = slide.shapes.add_picture(make_placeholder_image(), Inches(1), Inches(2), Inches(3), Inches(3))
    set_alt_text(picture, "A blue square used as a placeholder image")

    save_presentation(prs, "control_no_issues.pptx")


def make_issue_missing_alt_text():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    title_box.name = "Title 1"
    title_box.text_frame.text = "Slide With Missing Alt Text"

    picture = slide.shapes.add_picture(make_placeholder_image(), Inches(1), Inches(2), Inches(3), Inches(3))
    clear_alt_text(picture)

    save_presentation(prs, "issue_missing_alt_text.pptx")


def make_issue_no_title_shape():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    text_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    text_box.name = "TextBox 1"
    text_box.text_frame.text = "This slide has no shape named Title"

    save_presentation(prs, "issue_no_title_shape.pptx")


def make_issue_title_not_first_in_order():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    body_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    body_box.name = "Body Text 1"
    body_box.text_frame.text = "This shape was added first"

    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    title_box.name = "Title 1"
    title_box.text_frame.text = "This title was added second, so it is not first in reading order yet"

    save_presentation(prs, "issue_title_not_first_in_order.pptx")


def make_issue_empty_slide_no_shapes():
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])

    save_presentation(prs, "issue_empty_slide_no_shapes.pptx")


def make_issue_multiple_images_no_alt_text():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.3), Inches(8), Inches(0.7))
    title_box.name = "Title 1"
    title_box.text_frame.text = "Multiple Images Missing Alt Text"

    image_positions = [Inches(0.5), Inches(3), Inches(5.5)]
    for left in image_positions:
        picture = slide.shapes.add_picture(make_placeholder_image(), left, Inches(2), Inches(2), Inches(2))
        clear_alt_text(picture)

    save_presentation(prs, "issue_multiple_images_no_alt_text.pptx")


def make_issue_all_issues_combined():
    prs = Presentation()

    slide_one = prs.slides.add_slide(prs.slide_layouts[6])
    body_box = slide_one.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    body_box.name = "TextBox 1"
    body_box.text_frame.text = "No title shape on this slide, and this image has no alt text"
    picture = slide_one.shapes.add_picture(make_placeholder_image(), Inches(1), Inches(3.5), Inches(2), Inches(2))
    clear_alt_text(picture)

    prs.slides.add_slide(prs.slide_layouts[6])  # second slide is completely empty

    save_presentation(prs, "issue_all_issues_combined.pptx")


def make_issue_checks_sampler():
    prs = Presentation()

    first = prs.slides.add_slide(prs.slide_layouts[5])
    first.shapes.title.text = "Introduction"

    # the same title again, which is hard to tell apart when moving by heading
    second = prs.slides.add_slide(prs.slide_layouts[5])
    second.shapes.title.text = "Introduction"

    no_title = prs.slides.add_slide(prs.slide_layouts[6])
    body = no_title.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    body.name = "Body 1"
    body.text_frame.text = "Some content with no heading above it"

    link_slide = prs.slides.add_slide(prs.slide_layouts[5])
    link_slide.shapes.title.text = "Further reading"
    link_box = link_slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    link_run = link_box.text_frame.paragraphs[0].add_run()
    link_run.text = "click here"
    link_run.hyperlink.address = "https://example.com/syllabus"

    table_slide = prs.slides.add_slide(prs.slide_layouts[5])
    table_slide.shapes.title.text = "Results"
    table_shape = table_slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2))
    table_shape.table.first_row = False

    save_presentation(prs, "issue_checks_sampler.pptx")


def make_issue_every_check():
    prs = Presentation()

    # slide 1: enough writing for the reading level to be worked out
    intro = prs.slides.add_slide(prs.slide_layouts[5])
    intro.shapes.title.text = "Every Check On Its Own Slide"
    intro_box = intro.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
    intro_box.name = "Body 1"
    intro_box.text_frame.text = (
        "Each slide after this one is built to set off exactly one of the checks. "
        "The paragraph you are reading is only here so that there are enough words "
        "in the deck for the reading level to be worked out, because the measurement "
        "is skipped on presentations shorter than thirty words."
    )

    # slide 2: no title, because the only text shape is not named Title
    untitled = prs.slides.add_slide(prs.slide_layouts[6])
    untitled_box = untitled.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    untitled_box.name = "Body 1"
    untitled_box.text_frame.text = "There is no heading above this text."

    # slides 3 and 4: the same title twice
    first_results = prs.slides.add_slide(prs.slide_layouts[5])
    first_results.shapes.title.text = "Results"

    second_results = prs.slides.add_slide(prs.slide_layouts[5])
    second_results.shapes.title.text = "Results"

    # slide 5: a link whose words do not say where it goes
    link_slide = prs.slides.add_slide(prs.slide_layouts[5])
    link_slide.shapes.title.text = "Further Reading"
    link_box = link_slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    link_run = link_box.text_frame.paragraphs[0].add_run()
    link_run.text = "click here"
    link_run.hyperlink.address = "https://example.com/syllabus"

    # slide 6: a table with no header row, every cell filled so nothing else is flagged
    header_slide = prs.slides.add_slide(prs.slide_layouts[5])
    header_slide.shapes.title.text = "Marks By Term"
    header_table = header_slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    header_table.first_row = False
    filled_rows = [["Autumn", "72"], ["Winter", "78"], ["Spring", "81"]]
    for row_number in range(3):
        for column_number in range(2):
            header_table.cell(row_number, column_number).text = filled_rows[row_number][column_number]

    # slide 7: a table with a merged cell and an empty cell
    cells_slide = prs.slides.add_slide(prs.slide_layouts[5])
    cells_slide.shapes.title.text = "Timetable"
    cells_table = cells_slide.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(7), Inches(2)).table
    cells_table.cell(0, 0).text = "Day"
    cells_table.cell(0, 1).text = "Morning"
    cells_table.cell(0, 2).text = "Afternoon"
    cells_table.cell(1, 0).text = "Monday"
    cells_table.cell(1, 1).merge(cells_table.cell(1, 2))
    cells_table.cell(1, 1).text = "Lab session, runs all day"
    cells_table.cell(2, 0).text = "Tuesday"
    cells_table.cell(2, 1).text = "Lecture"
    # cell (2, 2) is left empty on purpose

    # slide 8: a shape parked past the right hand edge of the slide
    offslide = prs.slides.add_slide(prs.slide_layouts[5])
    offslide.shapes.title.text = "Notes To Myself"
    parked = offslide.shapes.add_textbox(Inches(11), Inches(2), Inches(3), Inches(1))
    parked.name = "Parked Note"
    parked.text_frame.text = "Old draft I meant to delete"

    # slide 9: the lower box was added first, so it is read before the one above it
    order_slide = prs.slides.add_slide(prs.slide_layouts[5])
    order_slide.shapes.title.text = "Method"
    lower_box = order_slide.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(1))
    lower_box.name = "Second Point"
    lower_box.text_frame.text = "Then we measured the current."
    upper_box = order_slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    upper_box.name = "First Point"
    upper_box.text_frame.text = "First we set up the circuit."

    # slide 10: something important left in the speaker notes
    notes_slide = prs.slides.add_slide(prs.slide_layouts[5])
    notes_slide.shapes.title.text = "Summary"
    notes_slide.notes_slide.notes_text_frame.text = (
        "Remember to tell them the exam covers chapters four and five."
    )

    # slide 11: nothing on it at all
    prs.slides.add_slide(prs.slide_layouts[6])

    save_presentation(prs, "issue_every_check.pptx")


def make_corrupt_empty_file():
    output_path = os.path.join(OUTPUT_FOLDER, "corrupt_empty_file.pptx")
    with open(output_path, "wb"):
        pass
    print("Created corrupt_empty_file.pptx")


def make_corrupt_random_bytes():
    output_path = os.path.join(OUTPUT_FOLDER, "corrupt_random_bytes.pptx")
    with open(output_path, "wb") as f:
        f.write(os.urandom(2000))
    print("Created corrupt_random_bytes.pptx")


def make_corrupt_plain_text_renamed():
    output_path = os.path.join(OUTPUT_FOLDER, "corrupt_plain_text_renamed.pptx")
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("This is just a plain text file pretending to be a PowerPoint file.\n")
    print("Created corrupt_plain_text_renamed.pptx")


def make_corrupt_zip_missing_pptx_parts():
    output_path = os.path.join(OUTPUT_FOLDER, "corrupt_zip_missing_pptx_parts.pptx")
    with zipfile.ZipFile(output_path, "w") as zip_file:
        zip_file.writestr("hello.txt", "This zip file is missing the parts a real pptx needs.")
    print("Created corrupt_zip_missing_pptx_parts.pptx")


def make_corrupt_truncated_valid_file():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    text_box.text_frame.text = "This file will be cut off partway through saving."

    full_file = io.BytesIO()
    prs.save(full_file)
    full_bytes = full_file.getvalue()
    truncated_bytes = full_bytes[:len(full_bytes) // 2]

    output_path = os.path.join(OUTPUT_FOLDER, "corrupt_truncated_valid_file.pptx")
    with open(output_path, "wb") as f:
        f.write(truncated_bytes)
    print("Created corrupt_truncated_valid_file.pptx")


if __name__ == "__main__":
    os.makedirs(OUTPUT_FOLDER, exist_ok=True)

    make_control_no_issues()
    make_issue_missing_alt_text()
    make_issue_no_title_shape()
    make_issue_title_not_first_in_order()
    make_issue_empty_slide_no_shapes()
    make_issue_multiple_images_no_alt_text()
    make_issue_all_issues_combined()
    make_issue_checks_sampler()
    make_issue_every_check()

    make_corrupt_empty_file()
    make_corrupt_random_bytes()
    make_corrupt_plain_text_renamed()
    make_corrupt_zip_missing_pptx_parts()
    make_corrupt_truncated_valid_file()

    print("Done. Files were written to: " + OUTPUT_FOLDER)
