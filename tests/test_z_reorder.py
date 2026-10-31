import os
import copy
import sys
import shutil

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# make sure the tests can import the project files from the folder above
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import title_detection
import z_reorder

FIXTURE_FOLDER = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Error Test PPTX")

# these fixture files should all fail to open
BROKEN_FILES = [
    "corrupt_empty_file.pptx",
    "corrupt_random_bytes.pptx",
    "corrupt_plain_text_renamed.pptx",
    "corrupt_zip_missing_pptx_parts.pptx",
    "corrupt_truncated_valid_file.pptx",
]

# these fixture files should all process successfully
WORKING_FILES = [
    "control_no_issues.pptx",
    "issue_missing_alt_text.pptx",
    "issue_no_title_shape.pptx",
    "issue_title_not_first_in_order.pptx",
    "issue_empty_slide_no_shapes.pptx",
    "issue_multiple_images_no_alt_text.pptx",
    "issue_all_issues_combined.pptx",
    "issue_low_contrast.pptx",
]


@pytest.fixture
def fake_captions(monkeypatch):
    def fake_caption(shape, slide_text=""):
        return "A test caption"

    monkeypatch.setattr(z_reorder, "generate_image_caption", fake_caption)


@pytest.fixture
def folders(tmp_path):
    input_folder = tmp_path / "uploads"
    output_folder = tmp_path / "processed"
    input_folder.mkdir()
    output_folder.mkdir()

    return str(input_folder), str(output_folder)


def copy_fixture(filename, input_folder):
    source = os.path.join(FIXTURE_FOLDER, filename)
    shutil.copy(source, os.path.join(input_folder, filename))


def test_tidy_caption_removes_filler_at_the_start():
    assert z_reorder.tidy_caption("The image shows a red car") == "A red car"


def test_tidy_caption_capitalises_the_first_letter():
    assert z_reorder.tidy_caption("there are two people talking") == "Two people talking"


def test_tidy_caption_leaves_a_normal_sentence_alone():
    assert z_reorder.tidy_caption("A bronze statue on a pedestal") == "A bronze statue on a pedestal"


def test_tidy_caption_handles_an_empty_caption():
    assert z_reorder.tidy_caption("") == ""


def test_build_caption_prompt_without_slide_text():
    prompt = z_reorder.build_caption_prompt("")

    assert prompt == z_reorder.CAPTION_PROMPT


def test_build_caption_prompt_includes_the_slide_text():
    prompt = z_reorder.build_caption_prompt("Root Causes of the Uprising")

    assert "Root Causes of the Uprising" in prompt


def test_get_slide_text_collects_the_words_on_the_slide():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "Soviet Response"

    assert "Soviet Response" in z_reorder.get_slide_text(slide)


def test_get_slide_text_is_cut_off_when_the_slide_is_very_long():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "word " * 500

    slide_text = z_reorder.get_slide_text(slide)

    assert len(slide_text) <= z_reorder.SLIDE_CONTEXT_MAX_CHARS


def test_get_slide_text_is_empty_for_a_slide_with_no_text():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    assert z_reorder.get_slide_text(slide) == ""


def get_first_picture(path):
    prs = Presentation(path)

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:
                return shape

    return None


def test_missing_alt_text_is_read_as_empty():
    picture = get_first_picture(os.path.join(FIXTURE_FOLDER, "issue_missing_alt_text.pptx"))

    assert not z_reorder.get_picture_alt_text(picture)


def test_existing_alt_text_is_read_back():
    picture = get_first_picture(os.path.join(FIXTURE_FOLDER, "control_no_issues.pptx"))

    assert z_reorder.get_picture_alt_text(picture)


def test_alt_text_can_be_set_and_read_again():
    picture = get_first_picture(os.path.join(FIXTURE_FOLDER, "issue_missing_alt_text.pptx"))

    z_reorder.set_picture_alt_text(picture, "A blue square")

    assert z_reorder.get_picture_alt_text(picture) == "A blue square"


def test_count_images_needing_captions():
    prs = Presentation(os.path.join(FIXTURE_FOLDER, "issue_multiple_images_no_alt_text.pptx"))

    assert z_reorder.count_images_needing_captions(prs) == 3


def test_pictures_that_already_have_alt_text_are_not_counted():
    prs = Presentation(os.path.join(FIXTURE_FOLDER, "control_no_issues.pptx"))

    assert z_reorder.count_images_needing_captions(prs) == 0


@pytest.mark.parametrize("filename", BROKEN_FILES)
def test_broken_files_fail_with_a_message_and_steps(filename, folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture(filename, input_folder)

    was_successful, message, steps = z_reorder.process_one_file(filename, input_folder, output_folder)

    assert was_successful is False
    assert message
    assert len(steps) > 0


def test_an_empty_file_says_it_is_empty(folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture("corrupt_empty_file.pptx", input_folder)

    was_successful, message, steps = z_reorder.process_one_file("corrupt_empty_file.pptx", input_folder, output_folder)

    assert was_successful is False
    assert "empty" in message.lower()


def test_a_missing_file_fails_politely(folders, fake_captions):
    input_folder, output_folder = folders

    was_successful, message, steps = z_reorder.process_one_file("not_here.pptx", input_folder, output_folder)

    assert was_successful is False
    assert message


@pytest.mark.parametrize("filename", WORKING_FILES)
def test_working_files_are_processed(filename, folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture(filename, input_folder)

    was_successful, message, steps = z_reorder.process_one_file(filename, input_folder, output_folder)

    assert was_successful is True
    assert message is None
    assert steps is None


def test_processing_writes_an_updated_file_and_a_report(folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture("issue_missing_alt_text.pptx", input_folder)

    z_reorder.process_one_file("issue_missing_alt_text.pptx", input_folder, output_folder)

    assert os.path.exists(os.path.join(output_folder, "issue_missing_alt_text_updated.pptx"))
    assert os.path.exists(os.path.join(output_folder, "issue_missing_alt_text_alt_text"))


def test_the_report_is_not_empty(folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture("issue_missing_alt_text.pptx", input_folder)

    z_reorder.process_one_file("issue_missing_alt_text.pptx", input_folder, output_folder)

    with open(os.path.join(output_folder, "issue_missing_alt_text_alt_text"), encoding="utf-8") as f:
        assert "Alt Text:" in f.read()


def test_the_score_counts_each_severity():
    summary = z_reorder.severity_summary(["critical", "serious", "serious", "minor"])

    assert "Accessibility score: 79/100" in summary
    assert "1 critical, 2 serious, 0 moderate, 1 minor" in summary


def test_the_score_never_goes_below_zero():
    assert "score: 0/100" in z_reorder.severity_summary(["critical"] * 20)


def test_the_report_opens_with_counts_by_severity(folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture("issue_missing_alt_text.pptx", input_folder)

    z_reorder.process_one_file("issue_missing_alt_text.pptx", input_folder, output_folder)

    with open(os.path.join(output_folder, "issue_missing_alt_text_alt_text"), encoding="utf-8") as f:
        report = f.read()

    assert report.startswith("Accessibility score: ")
    assert "Issues by severity: 1 critical" in report
    assert "[critical] The image had no real description" in report


def test_processing_logs_a_summary_with_the_output_paths(folders, fake_captions, caplog):
    input_folder, output_folder = folders
    copy_fixture("issue_missing_alt_text.pptx", input_folder)

    with caplog.at_level("INFO"):
        z_reorder.process_one_file("issue_missing_alt_text.pptx", input_folder, output_folder)

    summary = [m for m in caplog.messages if m.startswith("Summary:")]
    assert len(summary) == 1
    assert "image(s) captioned" in summary[0]
    assert "issue_missing_alt_text_updated.pptx" in summary[0]


def read_bytes(path):
    with open(path, "rb") as f:
        return f.read()


@pytest.mark.parametrize("filename", ["control_no_issues.pptx", "corrupt_truncated_valid_file.pptx"])
def test_the_original_is_left_unchanged_whether_processing_works_or_not(filename, folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture(filename, input_folder)
    original = os.path.join(input_folder, filename)
    before = read_bytes(original)

    z_reorder.process_one_file(filename, input_folder, output_folder)

    assert read_bytes(original) == before
    assert not os.path.exists(os.path.join(output_folder, filename))


def test_moving_originals_is_opt_in(folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture("control_no_issues.pptx", input_folder)

    z_reorder.process_one_file("control_no_issues.pptx", input_folder, output_folder, move_originals=True)

    assert not os.path.exists(os.path.join(input_folder, "control_no_issues.pptx"))
    assert os.path.exists(os.path.join(output_folder, "control_no_issues.pptx"))


def test_moving_an_original_never_replaces_a_file_already_in_the_output_folder(folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture("control_no_issues.pptx", input_folder)
    already_there = os.path.join(output_folder, "control_no_issues.pptx")

    with open(already_there, "wb") as f:
        f.write(b"someone else's file")

    was_successful, message, steps = z_reorder.process_one_file(
        "control_no_issues.pptx", input_folder, output_folder, move_originals=True)

    assert was_successful is False
    assert read_bytes(already_there) == b"someone else's file"
    assert os.path.exists(os.path.join(input_folder, "control_no_issues.pptx"))


def test_a_move_blocked_by_a_lock_finishes_on_the_next_try(folders, monkeypatch):
    input_folder, output_folder = folders
    copy_fixture("control_no_issues.pptx", input_folder)
    original = os.path.join(input_folder, "control_no_issues.pptx")
    before = read_bytes(original)

    real_remove = os.remove
    locked = True

    def remove(path):
        if locked:
            raise PermissionError("in use")
        real_remove(path)

    monkeypatch.setattr(z_reorder.os, "remove", remove)
    monkeypatch.setattr(z_reorder.time, "sleep", lambda seconds: None)

    assert z_reorder.move_original_to_output(input_folder, output_folder, "control_no_issues.pptx") is not None
    assert read_bytes(original) == before

    locked = False

    assert z_reorder.move_original_to_output(input_folder, output_folder, "control_no_issues.pptx") is None
    assert not os.path.exists(original)
    assert read_bytes(os.path.join(output_folder, "control_no_issues.pptx")) == before


def test_a_missing_description_is_written_into_the_updated_file(folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture("issue_missing_alt_text.pptx", input_folder)

    z_reorder.process_one_file("issue_missing_alt_text.pptx", input_folder, output_folder)

    picture = get_first_picture(os.path.join(output_folder, "issue_missing_alt_text_updated.pptx"))

    assert z_reorder.get_picture_alt_text(picture) == "A test caption"


def test_alt_text_that_was_already_there_is_kept(folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture("control_no_issues.pptx", input_folder)

    z_reorder.process_one_file("control_no_issues.pptx", input_folder, output_folder)

    picture = get_first_picture(os.path.join(output_folder, "control_no_issues_updated.pptx"))

    assert z_reorder.get_picture_alt_text(picture) != "A test caption"


def test_the_title_is_moved_to_the_front_of_the_slide(folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture("issue_title_not_first_in_order.pptx", input_folder)

    z_reorder.process_one_file("issue_title_not_first_in_order.pptx", input_folder, output_folder)

    prs = Presentation(os.path.join(output_folder, "issue_title_not_first_in_order_updated.pptx"))
    shape_names = []

    for shape in prs.slides[0].shapes:
        shape_names.append(shape.name)

    assert "Title" in shape_names[0]


def test_a_slide_with_no_shapes_does_not_stop_processing(folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture("issue_empty_slide_no_shapes.pptx", input_folder)

    was_successful, message, steps = z_reorder.process_one_file(
        "issue_empty_slide_no_shapes.pptx", input_folder, output_folder
    )

    assert was_successful is True


def test_a_caption_failure_falls_back_to_a_placeholder(folders, monkeypatch):
    input_folder, output_folder = folders
    copy_fixture("issue_missing_alt_text.pptx", input_folder)

    def broken_caption(shape, slide_text=""):
        raise RuntimeError("pretend the model failed")

    monkeypatch.setattr(z_reorder, "generate_image_caption", broken_caption)

    was_successful, message, steps = z_reorder.process_one_file(
        "issue_missing_alt_text.pptx", input_folder, output_folder
    )

    assert was_successful is True

    picture = get_first_picture(os.path.join(output_folder, "issue_missing_alt_text_updated.pptx"))

    assert "Image" in z_reorder.get_picture_alt_text(picture)


def test_progress_is_reported_for_every_image(folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture("issue_multiple_images_no_alt_text.pptx", input_folder)

    updates = []

    def remember_progress(done, total):
        updates.append((done, total))

    z_reorder.process_one_file(
        "issue_multiple_images_no_alt_text.pptx", input_folder, output_folder, remember_progress
    )

    # one update before starting, then one after each of the three images
    assert updates[0] == (0, 3)
    assert updates[-1] == (3, 3)


def test_a_ppt_file_is_rejected_politely_when_libreoffice_is_missing(folders, monkeypatch, fake_captions):
    input_folder, output_folder = folders

    # a .ppt cannot be opened directly, so a copied .pptx stands in for the upload
    shutil.copy(
        os.path.join(FIXTURE_FOLDER, "control_no_issues.pptx"),
        os.path.join(input_folder, "old_slides.ppt")
    )

    monkeypatch.setattr(z_reorder, "find_libreoffice", lambda: None)

    was_successful, message, steps = z_reorder.process_one_file("old_slides.ppt", input_folder, output_folder)

    assert was_successful is False
    assert ".ppt" in message
    assert len(steps) > 0


# --- title detection and reading order ---------------------------------------


def build_slide(names):
    """A blank slide holding plain text boxes, in the order given."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    for name in names:
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        box.name = name

    return slide


def names_in_order(slide):
    return [shape.name for shape in slide.shapes]


def test_a_title_that_is_not_first_is_moved_to_the_front():
    slide = build_slide(["Body 1", "Title 1"])

    assert z_reorder.fix_reading_order(slide) == 1
    assert names_in_order(slide) == ["Title 1", "Body 1"]


def test_two_titles_keep_their_relative_order():
    # the old code re-read slide.shapes[0] on every pass while reordering, so the
    # second title anchored against the first one after it had already moved and
    # these came back as Title B, Title A
    slide = build_slide(["Body 1", "Title A", "Title B"])

    assert z_reorder.fix_reading_order(slide) == 2
    assert names_in_order(slide) == ["Title A", "Title B", "Body 1"]


def test_a_title_already_at_the_front_is_left_alone():
    slide = build_slide(["Title 1", "Body 1"])

    assert z_reorder.fix_reading_order(slide) == 0
    assert names_in_order(slide) == ["Title 1", "Body 1"]


def test_reordering_a_second_time_changes_nothing():
    slide = build_slide(["Body 1", "Title 1"])

    z_reorder.fix_reading_order(slide)

    assert z_reorder.fix_reading_order(slide) == 0
    assert names_in_order(slide) == ["Title 1", "Body 1"]


def test_a_lowercase_title_name_is_still_a_title():
    slide = build_slide(["Body 1", "title 1"])

    assert z_reorder.fix_reading_order(slide) == 1
    assert names_in_order(slide) == ["title 1", "Body 1"]


def test_a_subtitle_is_not_treated_as_a_title():
    slide = build_slide(["Body 1", "Subtitle 2"])

    assert z_reorder.fix_reading_order(slide) == 0
    assert names_in_order(slide) == ["Body 1", "Subtitle 2"]


def test_a_slide_with_no_title_is_left_alone():
    slide = build_slide(["Body 1", "Body 2"])

    assert z_reorder.fix_reading_order(slide) == 0
    assert names_in_order(slide) == ["Body 1", "Body 2"]


def test_a_real_title_placeholder_is_found_even_when_renamed():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.name = "Header"

    assert z_reorder.is_title(title)


# --- titles that are not marked as titles --------------------------------------

BODY_TEXT = ("Evaporation, condensation and precipitation move water between the oceans, "
             "the air and the land, and the cycle repeats continuously.")


def build_text_slide(boxes):
    """A blank slide with plain text boxes given as (name, text, top in inches, font size), in z-order."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    for name, text, top, size in boxes:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(8.6), Inches(1))
        box.name = name
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(size)

    return slide


def test_a_title_typed_into_a_text_box_is_found_and_moved_to_the_front():
    # nothing here says "title": not the placeholder type, not the shape name
    slide = build_text_slide([("Body", BODY_TEXT, 2.5, 18), ("TextBox 7", "The Water Cycle", 0.4, 40)])

    assert z_reorder.fix_reading_order(slide) == 1
    assert names_in_order(slide) == ["TextBox 7", "Body"]


def test_a_title_below_a_small_kicker_line_is_still_found():
    slide = build_text_slide([("Body", BODY_TEXT, 3.0, 18),
                              ("Kicker", "CHAPTER 3", 0.2, 12),
                              ("Heading", "Evaporation and Rain", 0.7, 36)])

    z_reorder.fix_reading_order(slide)

    assert names_in_order(slide)[0] == "Heading"


def test_a_real_title_is_never_second_guessed(monkeypatch):
    monkeypatch.setattr(z_reorder, "guess_title", lambda slide: pytest.fail("guess_title was called"))
    slide = build_slide(["Body 1", "Title 1"])

    assert z_reorder.fix_reading_order(slide) == 1


def test_a_slide_of_only_body_paragraphs_gets_no_guessed_title():
    slide = build_text_slide([("A", BODY_TEXT, 1.0, 18), ("B", BODY_TEXT, 3.5, 18)])

    assert z_reorder.fix_reading_order(slide) == 0
    assert names_in_order(slide) == ["A", "B"]


def test_a_master_without_text_styles_does_not_break_font_size_lookup():
    # p:txStyles is optional; a placeholder with no size anywhere used to hit None.tag
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    master = slide.part.slide_layout.slide_master._element
    master.remove(master.find(title_detection.qn("p:txStyles")))

    assert title_detection.effective_pt(slide.shapes.title) == title_detection.DEFAULT_PT


def test_a_missing_title_model_leaves_the_slide_alone(monkeypatch):
    monkeypatch.setattr(title_detection, "_model", False)
    slide = build_text_slide([("Body", BODY_TEXT, 2.5, 18), ("TextBox 7", "The Water Cycle", 0.4, 40)])

    assert z_reorder.fix_reading_order(slide) == 0
    assert names_in_order(slide) == ["Body", "TextBox 7"]


# --- reading order beyond titles ----------------------------------------------


def add_footer(slide):
    """Copy the layout's footer placeholder onto the slide, the way PowerPoint does when footers are on."""
    layout_footer = next(ph for ph in slide.slide_layout.placeholders
                         if ph.placeholder_format.type == z_reorder.PP_PLACEHOLDER.FOOTER)
    slide.shapes._spTree.append(copy.deepcopy(layout_footer._element))

    return slide.shapes[-1]


def test_a_subtitle_read_before_the_title_is_moved_after_it():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title, subtitle = slide.shapes.title, slide.placeholders[1]
    title._element.addprevious(subtitle._element)

    assert z_reorder.fix_reading_order(slide) == 1
    assert names_in_order(slide) == [title.name, subtitle.name]


def test_a_footer_read_before_the_content_is_moved_to_the_end():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Results"
    footer = add_footer(slide)
    body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    slide.shapes.title._element.addnext(footer._element)

    assert z_reorder.fix_reading_order(slide) == 1
    assert names_in_order(slide) == [slide.shapes.title.name, body.name, footer.name]
    assert z_reorder.fix_reading_order(slide) == 0


def test_content_shapes_keep_the_order_the_author_gave_them():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Results"
    diagram = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(2), Inches(2))
    body = slide.shapes.add_textbox(Inches(4), Inches(2), Inches(4), Inches(1))
    slide.shapes.title._element.addprevious(body._element)

    z_reorder.fix_reading_order(slide)

    assert names_in_order(slide) == [slide.shapes.title.name, body.name, diagram.name]


def build_positioned_slide(boxes):
    """A blank slide with text boxes given as (name, left, top) in inches, in z-order."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    for name, left, top in boxes:
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(3), Inches(1))
        box.name = name

    return slide


def test_a_shape_above_another_but_read_after_it_is_found():
    slide = build_positioned_slide([("Lower", 1, 4), ("Upper", 1, 2)])

    assert z_reorder.find_out_of_visual_order(slide) == [("Upper", "Lower")]


def test_shapes_read_top_to_bottom_are_not_reported():
    slide = build_positioned_slide([("Upper", 1, 2), ("Lower", 1, 4)])

    assert z_reorder.find_out_of_visual_order(slide) == []


def test_shapes_in_different_columns_are_not_compared():
    # left column read first, even though the right column starts higher up
    slide = build_positioned_slide([("Left", 1, 4), ("Right", 5, 2)])

    assert z_reorder.find_out_of_visual_order(slide) == []


def test_shapes_read_out_of_visual_order_are_listed_in_the_report(folders, fake_captions):
    slide = build_positioned_slide([("Lower", 1, 4), ("Upper", 1, 2)])

    ok, _, report, updated = run_pipeline_on(slide.part.package.presentation_part.presentation, folders)

    assert ok
    assert "Slide 1, Shape: Upper is above Lower but is read after it" in report
    assert names_in_order(updated.slides[0]) == ["Lower", "Upper"]


# --- alt text that is not really alt text ------------------------------------


@pytest.mark.parametrize("alt_text", [
    "diagram.png",
    "Screenshot 2024.PNG",
    "chart.jpeg",
    " figure.svg ",
])
def test_filename_style_alt_text_is_not_a_description(alt_text):
    assert z_reorder.is_placeholder_alt_text(alt_text, "Picture 3")


@pytest.mark.parametrize("name", [
    "Picture 3",
    "Image 5",
    "Graphic 7",
    "(Picture 12)",
    "Picture Placeholder 2",
    "Content Placeholder 4",
])
def test_alt_text_repeating_an_auto_generated_shape_name_is_not_a_description(name):
    assert z_reorder.is_placeholder_alt_text(name, name)


@pytest.mark.parametrize("name", ["(Picture 3", "Picture 3)"])
def test_an_unbalanced_parenthesis_is_not_an_auto_generated_name(name):
    assert not z_reorder.is_placeholder_alt_text(name, name)


@pytest.mark.parametrize("name", [
    "Water cycle diagram",
    "Free body diagram of the beam",
    "Beam",
])
def test_a_description_matching_an_author_chosen_shape_name_is_kept(name):
    # someone who renames a shape and writes the same text as alt text meant it.
    # Replacing real alt text is worse than leaving a thin description alone.
    assert not z_reorder.is_placeholder_alt_text(name, name)


def test_a_real_description_on_an_auto_named_shape_is_kept():
    assert not z_reorder.is_placeholder_alt_text("A blue square", "Picture 3")


@pytest.mark.parametrize("alt_text", [
    "A bar chart of enrolment by year",
    "",
    None,
])
def test_a_real_description_is_left_alone(alt_text):
    assert not z_reorder.is_placeholder_alt_text(alt_text, "Picture 3")


def set_alt_text_on_every_picture(prs, alt_text):
    for slide in prs.slides:
        for shape in slide.shapes:
            if z_reorder.is_picture(shape):
                z_reorder.set_picture_alt_text(shape, alt_text)


def test_a_picture_whose_alt_text_is_a_filename_still_needs_a_caption():
    prs = Presentation(os.path.join(FIXTURE_FOLDER, "issue_missing_alt_text.pptx"))
    set_alt_text_on_every_picture(prs, "diagram.png")

    assert z_reorder.count_images_needing_captions(prs) == 1


def test_a_picture_with_a_real_description_does_not_need_a_caption():
    prs = Presentation(os.path.join(FIXTURE_FOLDER, "issue_missing_alt_text.pptx"))
    set_alt_text_on_every_picture(prs, "A blue square on a white background")

    assert z_reorder.count_images_needing_captions(prs) == 0


# --- shape types python-pptx cannot resolve ----------------------------------


def test_a_text_box_is_not_a_picture():
    slide = build_slide(["Body 1"])

    assert not z_reorder.is_picture(slide.shapes[0])


def test_an_unresolvable_shape_type_does_not_raise():
    class Unresolvable:
        name = "Odd 1"

        @property
        def shape_type(self):
            raise NotImplementedError

    assert z_reorder.get_shape_type(Unresolvable()) is None
    assert not z_reorder.is_picture(Unresolvable())
    assert not z_reorder.needs_caption(Unresolvable())


def test_a_whole_folder_is_processed_in_one_run(folders, fake_captions):
    input_folder, output_folder = folders
    for filename in WORKING_FILES + BROKEN_FILES:
        copy_fixture(filename, input_folder)

    z_reorder.accessibility_processor(input_folder, output_folder)

    for filename in WORKING_FILES:
        assert os.path.exists(os.path.join(output_folder, filename.replace(".pptx", "_updated.pptx")))

    for filename in BROKEN_FILES:
        assert not os.path.exists(os.path.join(output_folder, filename.replace(".pptx", "_updated.pptx")))

    prs = Presentation(os.path.join(output_folder, "issue_title_not_first_in_order_updated.pptx"))
    assert [shape.name for shape in prs.slides[0].shapes] == ["Title 1", "Body Text 1"]


def make_deck_with_untitled_slides(path):
    prs = Presentation()
    titled = prs.slides.add_slide(prs.slide_layouts[5])
    titled.shapes.title.text = "Has a title"
    prs.slides.add_slide(prs.slide_layouts[5])  # title placeholder left empty
    prs.slides.add_slide(prs.slide_layouts[6])  # blank layout, no shapes
    prs.save(path)


def test_slides_without_a_title_are_listed_in_the_report(folders, fake_captions):
    input_folder, output_folder = folders
    make_deck_with_untitled_slides(os.path.join(input_folder, "untitled.pptx"))

    z_reorder.process_one_file("untitled.pptx", input_folder, output_folder)

    with open(os.path.join(output_folder, "untitled_alt_text"), encoding="utf-8") as f:
        report = f.read()

    assert "Slides with no title (WCAG 2.4.2): 2, 3" in report


def test_a_deck_where_every_slide_has_a_title_reports_nothing(folders, fake_captions):
    input_folder, output_folder = folders
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Only slide"
    prs.save(os.path.join(input_folder, "titled.pptx"))

    z_reorder.process_one_file("titled.pptx", input_folder, output_folder)

    with open(os.path.join(output_folder, "titled_alt_text"), encoding="utf-8") as f:
        assert "no title" not in f.read()


def run_pipeline_on(prs, folders, language=None):
    """Save prs as a deck, process it, and return the report text and the updated deck."""
    input_folder, output_folder = folders
    prs.save(os.path.join(input_folder, "deck.pptx"))

    ok, message, _ = z_reorder.process_one_file("deck.pptx", input_folder, output_folder, language=language)

    with open(os.path.join(output_folder, "deck_alt_text"), encoding="utf-8") as f:
        report = f.read()

    return ok, message, report, Presentation(os.path.join(output_folder, "deck_updated.pptx"))


def make_language_deck(run_languages):
    """One text box per entry, each run marked with that lang value (None leaves it unmarked)."""
    prs = Presentation()

    for language in run_languages:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "A title"
        run = slide.shapes.title.text_frame.paragraphs[0].runs[0]

        if language is not None:
            run._r.get_or_add_rPr().set("lang", language)

    return prs


def remove_default_language(prs):
    default_style = prs.part._element.find(qn("p:defaultTextStyle"))
    default_style.getparent().remove(default_style)


def run_languages_of(prs):
    return {run_props.get("lang") for slide in prs.slides for run_props in slide._element.iter(qn("a:rPr"))}


def default_language_of(prs):
    return z_reorder.get_default_language(prs)


def test_a_deck_with_one_valid_language_reports_nothing_about_language(folders, fake_captions):
    _, _, report, _ = run_pipeline_on(make_language_deck(["en-US", "en-US"]), folders)

    assert "Language" not in report


def test_regional_variants_of_one_language_are_not_mixed(folders, fake_captions):
    _, _, report, _ = run_pipeline_on(make_language_deck(["en-US", "en-GB"]), folders)

    assert "Language" not in report


def test_a_deck_with_no_default_language_is_reported(folders, fake_captions):
    prs = make_language_deck([None])
    remove_default_language(prs)

    _, _, report, _ = run_pipeline_on(prs, folders)

    assert "Language (WCAG 3.1.1, 3.1.2): the presentation has no default language" in report


def test_an_invalid_language_value_is_reported_with_its_slides(folders, fake_captions):
    _, _, report, _ = run_pipeline_on(make_language_deck(["en-US", "english", "en-US", "english"]), folders)

    assert "'english' is not a valid language code (slides 2, 4)" in report


def test_mixed_languages_are_reported_with_the_slides_of_the_minority(folders, fake_captions):
    _, _, report, _ = run_pipeline_on(make_language_deck(["en-US", "en-US", "fr-FR"]), folders)

    assert "more than one language is used" in report
    assert "en-US on 2 slide(s)" in report
    assert "fr-FR on slide(s) 3" in report


def test_reporting_alone_does_not_change_the_languages(folders, fake_captions):
    prs = make_language_deck(["en-US", "fr-FR"])

    _, _, _, updated = run_pipeline_on(prs, folders)

    assert run_languages_of(updated) == {"en-US", "fr-FR"}


def test_a_chosen_language_is_set_across_the_deck(folders, fake_captions):
    prs = make_language_deck(["en-US", "fr-FR", "english", None])

    ok, _, report, updated = run_pipeline_on(prs, folders, language="es-ES")

    assert ok
    assert default_language_of(updated) == "es-ES"
    assert run_languages_of(updated) - {None} == {"es-ES"}
    assert "Language set to es-ES" in report


def test_setting_a_language_also_fixes_the_masters_and_layouts(folders, fake_captions):
    _, _, _, updated = run_pipeline_on(make_language_deck(["en-US"]), folders, language="ja-JP")

    for master in updated.slide_masters:
        langs = {e.get("lang") for e in master._element.iter(qn("a:defRPr"), qn("a:rPr")) if e.get("lang")}
        assert langs <= {"ja-JP"}


def test_a_deck_with_no_default_style_gets_one_when_a_language_is_set(folders, fake_captions):
    prs = make_language_deck([None])
    remove_default_language(prs)

    _, _, _, updated = run_pipeline_on(prs, folders, language="de-DE")

    assert default_language_of(updated) == "de-DE"


def test_a_deck_with_a_missing_language_is_fixed_and_no_longer_reported(folders, fake_captions):
    prs = make_language_deck([None])
    remove_default_language(prs)
    _, _, _, updated = run_pipeline_on(prs, folders, language="en-US")

    assert z_reorder.find_language_problems(updated) == []


def test_a_language_that_is_not_a_code_is_refused(folders, fake_captions):
    input_folder, output_folder = folders
    make_language_deck(["en-US"]).save(os.path.join(input_folder, "deck.pptx"))

    ok, message, steps = z_reorder.process_one_file("deck.pptx", input_folder, output_folder, language="English")

    assert not ok
    assert "not a valid language code" in message
    assert steps



def add_sized_box(slide, text, size):
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)

    return box


def small_text_names(slide, **kwargs):
    return [name for name, _, _ in z_reorder.find_small_text(slide, **kwargs)]


def test_an_explicit_size_below_the_minimum_is_flagged():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    small = add_sized_box(slide, "Footnote", 8)
    add_sized_box(slide, "Body", 9)

    assert z_reorder.find_small_text(slide) == [(small.name, 8, "Footnote")]


def test_a_size_inherited_from_the_master_is_flagged():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.placeholders[1].text = "Body text"
    master = slide.slide_layout.slide_master._element
    body_style = master.find(title_detection.qn("p:txStyles")).find(title_detection.qn("p:bodyStyle"))
    body_style.find(title_detection.qn("a:lvl1pPr")).find(title_detection.qn("a:defRPr")).set("sz", "800")

    assert small_text_names(slide) == [slide.placeholders[1].name]


def test_a_size_inherited_from_the_layout_wins_over_the_master():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.placeholders[1].text = "Body text"
    layout_body = [p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == 1][0]
    list_style = layout_body._element.find(".//" + title_detection.qn("a:lstStyle"))
    level_one = etree.SubElement(list_style, title_detection.qn("a:lvl1pPr"))
    etree.SubElement(level_one, title_detection.qn("a:defRPr")).set("sz", "700")

    assert small_text_names(slide) == [slide.placeholders[1].name]


def test_inherited_sizes_at_normal_size_are_not_flagged():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Title"
    slide.placeholders[1].text = "Body text"

    assert z_reorder.find_small_text(slide) == []


def test_a_shrink_on_overflow_scale_can_push_text_below_the_minimum():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = add_sized_box(slide, "Squeezed", 12)
    body_pr = box.text_frame._txBody.find(title_detection.qn("a:bodyPr"))
    fit = etree.SubElement(body_pr, title_detection.qn("a:normAutofit"))
    fit.set("fontScale", "60000")

    assert small_text_names(slide) == [box.name]


def test_the_minimum_can_be_changed():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = add_sized_box(slide, "Twelve point", 12)

    assert small_text_names(slide) == []
    assert small_text_names(slide, min_pt=14) == [box.name]


def test_small_text_in_a_table_cell_is_flagged():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(4), Inches(1))
    run = table.table.cell(0, 0).text_frame.paragraphs[0].add_run()
    run.text = "Cell"
    run.font.size = Pt(7)

    assert small_text_names(slide) == [table.name]


def test_small_text_is_listed_by_slide_in_the_report(folders, fake_captions):
    input_folder, output_folder = folders
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    add_sized_box(prs.slides.add_slide(prs.slide_layouts[6]), "Tiny source note", 8)
    prs.save(os.path.join(input_folder, "small.pptx"))

    z_reorder.process_one_file("small.pptx", input_folder, output_folder)

    with open(os.path.join(output_folder, "small_alt_text"), encoding="utf-8") as f:
        report = f.read()

    assert "Text below 9 pt (WCAG 1.4.4)" in report
    assert "Slide 2, Shape: TextBox" in report
    assert "8 pt: 'Tiny source note'" in report
    assert "Slide 1," not in report


def add_colored_box(slide, text, hex_color=None, size=14):
    box = add_sized_box(slide, text, size)

    if hex_color:
        box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(hex_color)

    return box


def contrast_names(slide):
    return [(name, needed) for name, _, needed, _ in z_reorder.find_low_contrast(slide)]


def test_contrast_ratio_matches_wcag():
    assert round(z_reorder.contrast_ratio((0, 0, 0), (1, 1, 1)), 2) == 21
    assert round(z_reorder.contrast_ratio((1, 1, 1), (1, 1, 1)), 2) == 1
    grey = tuple(0x77 / 255 for _ in range(3))
    assert round(z_reorder.contrast_ratio(grey, (1, 1, 1)), 2) == 4.48


def test_default_text_on_the_default_background_passes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Title"
    slide.placeholders[1].text = "Body text"
    add_colored_box(slide, "Plain box")

    assert z_reorder.find_low_contrast(slide) == []


def test_light_text_on_white_fails():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    faint = add_colored_box(slide, "Faint", "AAAAAA")

    [(name, ratio, needed, text)] = z_reorder.find_low_contrast(slide)
    assert (name, round(ratio, 2), needed, text) == (faint.name, 2.32, 4.5, "Faint")


def test_large_text_needs_only_three_to_one():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_colored_box(slide, "Large", "949494", size=18)
    bold = add_colored_box(slide, "Large bold", "949494", size=14)
    bold.text_frame.paragraphs[0].runs[0].font.bold = True
    small = add_colored_box(slide, "Small", "949494", size=14)

    assert contrast_names(slide) == [(small.name, 4.5)]


def test_a_theme_colour_on_a_solid_slide_background_is_resolved():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = add_sized_box(slide, "White text", 14)
    rpr = box.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:schemeClr")).set("val", "bg1")

    assert contrast_names(slide) == [(box.name, 4.5)]

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0x1F, 0x1F, 0x1F)

    assert contrast_names(slide) == []


def test_lum_mod_darkens_a_theme_colour():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = add_sized_box(slide, "Darkened white", 14)
    rpr = box.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    color = etree.SubElement(etree.SubElement(rpr, qn("a:solidFill")), qn("a:schemeClr"))
    color.set("val", "bg1")
    etree.SubElement(color, qn("a:lumMod")).set("val", "25000")

    assert contrast_names(slide) == []


def test_a_filled_shape_behind_the_text_is_its_background():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(8), Inches(4))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(0, 0, 0)
    on_panel = add_colored_box(slide, "White on black", "FFFFFF")

    assert contrast_names(slide) == []

    on_panel.fill.solid()
    on_panel.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xF0)

    assert contrast_names(slide) == [(on_panel.name, 4.5)]


def test_a_gradient_background_is_listed_as_not_checked():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.gradient()
    box = add_colored_box(slide, "On a gradient", "000000")

    assert z_reorder.find_low_contrast(slide) == [(box.name, None, None, "On a gradient")]


def test_low_contrast_is_listed_by_slide_in_the_report(folders, fake_captions):
    input_folder, output_folder = folders
    copy_fixture("issue_low_contrast.pptx", input_folder)

    z_reorder.process_one_file("issue_low_contrast.pptx", input_folder, output_folder)

    with open(os.path.join(output_folder, "issue_low_contrast_alt_text"), encoding="utf-8") as f:
        report = f.read()

    assert "Text with too little contrast (WCAG 1.4.3)" in report
    assert "Slide 2, Shape: Faint Note - 2.32:1, needs 4.5:1: 'Light grey on white'" in report
    assert "Slide 2, Shape: Grey Body - 3.03:1, needs 4.5:1" in report
    assert "Slide 1," not in report
    assert "Contrast not checked" not in report


def test_a_negative_lum_off_stops_at_black():
    color = etree.Element(qn("a:srgbClr"), val="000000")
    etree.SubElement(color, qn("a:lumOff")).set("val", "-50000")

    assert z_reorder.color_to_rgb(color, None, None) == (0.0, 0.0, 0.0)


def add_link(slide, parts, address="https://www.washington.edu/accessibility/"):
    """A text box whose runs are the given parts, all linked to address."""
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragraph = box.text_frame.paragraphs[0]

    for part in parts:
        run = paragraph.add_run()
        run.text = part
        run.hyperlink.address = address

    return box


@pytest.mark.parametrize("text", [
    "click here", "Click Here!", "here", "link", "this link", "Read more...", "Learn more",
    "more", "more info", "website", "  ", "https://www.washington.edu", "www.washington.edu/accessibility",
])
def test_vague_link_text_is_flagged(text):
    assert z_reorder.is_vague_link_text(text)


@pytest.mark.parametrize("text", [
    "UW accessibility guidelines", "Download the syllabus (PDF)", "Here is the course calendar", "uw.edu",
])
def test_descriptive_link_text_is_not_flagged(text):
    assert not z_reorder.is_vague_link_text(text)


def test_a_link_split_over_runs_is_read_as_one_link():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    vague = add_link(slide, ["click ", "here"])
    add_link(slide, ["UW ", "accessibility guidelines"])

    assert z_reorder.find_vague_links(slide) == [
        (vague.name, "click here", "https://www.washington.edu/accessibility/")]


def test_unlinked_vague_words_are_not_flagged():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_sized_box(slide, "Click here", 14)

    assert z_reorder.find_vague_links(slide) == []


def test_a_linked_shape_is_flagged_by_its_text_or_alt_text():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    button = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(2), Inches(1))
    button.text_frame.text = "More"
    button.click_action.hyperlink.address = "https://example.com"
    icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3), Inches(0), Inches(1), Inches(1))
    icon._element._nvXxPr.cNvPr.set("descr", "Course website")
    icon.click_action.hyperlink.address = "https://example.com/course"

    assert z_reorder.find_vague_links(slide) == [(button.name, "More", "https://example.com")]


def test_a_linked_picture_is_flagged_by_its_alt_text():
    picture = get_first_picture(os.path.join(FIXTURE_FOLDER, "issue_missing_alt_text.pptx"))
    z_reorder.set_picture_alt_text(picture, "Click here")
    picture.click_action.hyperlink.address = "https://example.com"

    assert z_reorder.find_vague_links(picture.part.slide) == [(picture.name, "Click here", "https://example.com")]


def test_vague_links_are_listed_by_slide_in_the_report(folders, fake_captions):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "No links"
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Resources"
    add_link(slide, ["here"], "https://example.com/syllabus.pdf")

    _, _, report, _ = run_pipeline_on(prs, folders)

    assert "[moderate] Link text does not say where the link goes (WCAG 2.4.4)" in report
    assert "Slide 2, Shape: TextBox 2 - 'here' links to https://example.com/syllabus.pdf" in report
    assert "Slide 1," not in report
