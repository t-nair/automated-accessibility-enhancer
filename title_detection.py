"""Find the title of a slide that has no real title placeholder.

PowerPoint decks often hold the title in a plain text box or a body placeholder, which
the placeholder/name rule in z_reorder cannot see. This module turns each slide into one
record per text shape (geometry, effective font size, formatting, wording) and asks a small
gradient-boosted model which shape, if any, is the title. The model is trained by
title_research/train_final.py; see title_research/README.md for how it was chosen.
"""
import logging
import re
from pathlib import Path

import numpy as np
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

TITLE_PH = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
FURNITURE_PH = (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)
DEFAULT_PT = 18.0


def _first_sz(el):
    """First explicit font size (points) under an lstStyle/defRPr style element."""
    if el is None:
        return None
    for node in el.iter(qn("a:defRPr"), qn("a:rPr")):
        if node.get("sz"):
            return int(node.get("sz")) / 100
    return None


def _ph_chain(shape):
    """The same placeholder on the layout and on the master, if they exist."""
    chain = []
    try:
        layout = shape.part.slide_layout
        idx = shape.placeholder_format.idx
        for ph in layout.placeholders:
            if ph.placeholder_format.idx == idx:
                chain.append(ph._element)
        master = layout.slide_master
        for ph in master.placeholders:
            if ph.placeholder_format.type == shape.placeholder_format.type:
                chain.append(ph._element)
        tx_styles = master._element.find(qn("p:txStyles"))
        if tx_styles is not None:
            chain.append(tx_styles)
    except Exception:
        pass
    return chain


def effective_pt(shape):
    """Largest font size the shape's text will render at, in points."""
    tf = shape.text_frame
    sizes = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size and r.text.strip()]
    scale = 1.0
    body = tf._txBody.find(qn("a:bodyPr"))
    fit = body.find(qn("a:normAutofit")) if body is not None else None
    if fit is not None and fit.get("fontScale"):
        scale = int(fit.get("fontScale")) / 100000
    if sizes:
        return max(sizes) * scale
    size = _first_sz(tf._txBody.find(qn("a:lstStyle")))
    if size is None and shape.is_placeholder:
        is_title = shape.placeholder_format.type in TITLE_PH
        for el in _ph_chain(shape):
            if el.tag == qn("p:txStyles"):
                style = el.find(qn("p:titleStyle" if is_title else "p:bodyStyle"))
                size = _first_sz(style)
            else:
                size = _first_sz(el.find(".//" + qn("a:lstStyle")))
            if size:
                break
    return (size or DEFAULT_PT) * scale


def _is_bold(shape):
    runs = [r for p in shape.text_frame.paragraphs for r in p.runs if r.text.strip()]
    return bool(runs) and sum(bool(r.font.bold) for r in runs) * 2 >= len(runs)


def _alignment(shape):
    for p in shape.text_frame.paragraphs:
        if p.alignment is not None:
            return str(p.alignment).split(".")[-1].split(" ")[0]
    return "none"


def _has_bullet(shape):
    for p in shape.text_frame.paragraphs:
        ppr = p._p.pPr
        if p.level > 0 or (ppr is not None and ppr.find(qn("a:buChar")) is not None):
            return True
    return False


def _safe_type(shape):
    try:
        return shape.shape_type
    except Exception:
        return None


def _walk(shapes, tx=(0, 0, 1, 1)):
    """Yield (shape, abs_left, abs_top, abs_w, abs_h) in EMU, descending into groups."""
    ox, oy, sx, sy = tx
    for sh in shapes:
        try:
            left, top, w, h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            left = top = w = h = None
        if left is None:
            left = top = w = h = 0
        if _safe_type(sh) == MSO_SHAPE_TYPE.GROUP:
            x = sh._element.grpSpPr.find(qn("a:xfrm"))
            ch = x.find(qn("a:chOff")), x.find(qn("a:chExt"))
            cox, coy = int(ch[0].get("x")), int(ch[0].get("y"))
            cw, chh = int(ch[1].get("cx")) or 1, int(ch[1].get("cy")) or 1
            ax, ay, aw, ah = ox + left * sx, oy + top * sy, w * sx, h * sy
            yield from _walk(sh.shapes, (ax - cox * aw / cw, ay - coy * ah / chh, aw / cw, ah / chh))
        else:
            yield sh, ox + left * sx, oy + top * sy, w * sx, h * sy


def slide_records(slide, slide_no, n_slides, slide_w, slide_h):
    """One dict per non-empty text shape on the slide, in z-order, plus slide context."""
    recs = []
    for z, (sh, left, top, w, h) in enumerate(_walk(slide.shapes)):
        if not getattr(sh, "has_text_frame", False) or not sh.text_frame.text.strip():
            continue
        text = " ".join(sh.text_frame.text.split())
        ph = None
        if sh.is_placeholder:
            try:
                ph = sh.placeholder_format.type
            except Exception:
                ph = None
        paras = [p for p in sh.text_frame.paragraphs if p.text.strip()]
        r = {
            "shape": sh, "z": z, "text": text,
            "n_chars": len(text), "n_words": len(text.split()), "n_paras": len(paras),
            "x": left / slide_w, "y": top / slide_h, "w": w / slide_w, "h": h / slide_h,
            "pt": effective_pt(sh), "bold": _is_bold(sh), "align": _alignment(sh),
            "bullets": _has_bullet(sh), "caps": text.isupper() and any(c.isalpha() for c in text),
            "end_punct": text[-1] in ".,;:!?" if text else False,
            "digits_only": bool(re.fullmatch(r"[\d\s/.\-]+", text)),
            "rot": getattr(sh, "rotation", 0) or 0,
            "is_auto": _safe_type(sh) == MSO_SHAPE_TYPE.AUTO_SHAPE,
            "gold": ph in TITLE_PH,
        }
        r["cx"], r["cy"] = r["x"] + r["w"] / 2, r["y"] + r["h"] / 2
        # structural cues: the model never sees them, the baseline rule and the footer filter do
        r["ph"] = str(ph).split(".")[-1].split(" ")[0] if ph is not None else None
        r["name"] = sh.name
        r["furniture"] = ph in FURNITURE_PH
        recs.append(r)
    ctx = {"slide_no": slide_no, "n_slides": n_slides,
           "first": slide_no == 1, "last": slide_no == n_slides}
    for r in recs:
        r.update(ctx)
    return add_context(recs)


def add_context(recs):
    """Slide-relative features; call again whenever shapes are added or moved."""
    if recs:
        mx = max(r["pt"] for r in recs)
        order = sorted(r["pt"] for r in recs)
        ys = sorted(r["y"] for r in recs)
        for r in recs:
            r["cx"], r["cy"] = r["x"] + r["w"] / 2, r["y"] + r["h"] / 2
            r["pt_rel"] = r["pt"] / mx
            r["pt_rank"] = sum(p > r["pt"] for p in order)           # shapes with a bigger font
            r["y_rank"] = sum(y < r["y"] for y in ys)                # shapes strictly higher up
            r["n_shapes"] = len(recs)
    return recs


def deck_slides(path):
    """Yield (slide_no, records) for every slide; gold = the real title placeholder."""
    from pptx import Presentation
    prs = Presentation(path)
    n = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        yield i, slide_records(slide, i, n, prs.slide_width or 1, prs.slide_height or 1)


# ---------------------------------------------------------------- the classifier
ALIGNS = ("center", "left", "right", "none")
NUM = ["x", "y", "w", "h", "cx", "cy", "pt", "pt_rel", "pt_rank", "y_rank", "n_chars", "n_words",
       "n_paras", "rot", "n_shapes"]
BOOL = ["bold", "bullets", "caps", "end_punct", "digits_only", "is_auto", "first", "last"]
# z: where the shape sits in z-order. Real decks mostly put the title first (argmax accuracy on
# hand-labeled slides 96% with it, 92% without); the z_shuffle perturbation stops the model
# from relying on it when the order is scrambled.
# text: cheap wording features (title case, "?", numbered heading ...).
CONFIG = {"z": True, "text": True}

MODEL_PATH = Path(__file__).with_name("title_model.joblib")
# below this the best shape is not called a title. Overall accuracy is flat for 0.05-0.2 on the
# hand-labeled slides; lower = fewer missed titles, higher = fewer titles invented on diagram slides.
THRESHOLD = 0.1


def text_feats(t):
    words = t.split()
    caps = sum(w[:1].isupper() for w in words) / len(words)
    return [caps, float(t.endswith("?")), float(bool(re.match(r"^\(?\d+[.)]?\s", t))), float(":" in t),
            sum(map(len, words)) / len(words), float(t.count(",") + t.count(";")), float(t[:1].islower())]


def feature_rows(recs):
    """Numeric matrix, one row per shape. Placeholder type and shape name are never used."""
    max_chars = max(r["n_chars"] for r in recs)
    rows = []
    for i, r in enumerate(recs):
        row = [float(r[k]) for k in NUM] + [float(r[k]) for k in BOOL]
        row += [float(r["align"].upper().startswith(a.upper())) for a in ALIGNS]
        row += [r["y"] + r["h"], r["w"] * r["h"], r["n_chars"] / max_chars,
                r["n_chars"] / max(r["n_paras"], 1), float(abs(r["cx"] - 0.5) < 0.08)]
        if CONFIG["z"]:
            row.append(i / len(recs))
        if CONFIG["text"]:
            row += text_feats(r["text"])
        rows.append(row)
    return np.array(rows)


_model = None


def get_model():
    """Loaded on first use. Returns None (title guessing off) if the file or scikit-learn is unusable."""
    global _model
    if _model is None:
        try:
            import joblib
            saved = joblib.load(MODEL_PATH)
            if saved["config"] != CONFIG:
                raise ValueError(f"model was trained with {saved['config']}, code uses {CONFIG}")
            _model = saved["model"]
        except Exception as e:
            logging.warning(f"Title model unavailable, only real title placeholders will be found. Error: {e}")
            _model = False
    return _model or None


def guess_title(slide):
    """The text shape most likely to be this slide's title, or None.

    Only meant for slides without a real title placeholder. Footer, date and slide-number
    placeholders are never candidates. Never raises: a slide it cannot read has no guess.
    """
    model = get_model()
    if model is None:
        return None
    try:
        prs = slide.part.package.presentation_part.presentation
        ids = [s.slide_id for s in prs.slides]
        recs = slide_records(slide, ids.index(slide.slide_id) + 1, len(ids), prs.slide_width or 1, prs.slide_height or 1)
        recs = add_context([r for r in recs if not r["furniture"]])
        if not recs:
            return None
        p = model.predict_proba(feature_rows(recs))[:, 1]
        best = int(np.argmax(p))
        return recs[best]["shape"] if p[best] >= THRESHOLD else None
    except Exception as e:
        logging.warning(f"Could not guess a title for a slide. Error: {e}")
        return None
